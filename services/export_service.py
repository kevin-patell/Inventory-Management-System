import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import Config

def export_to_csv(headers, rows):
    """
    Generates a CSV string/bytes from headers and rows.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return output.getvalue().encode('utf-8-sig')

def export_to_excel(title, headers, rows):
    """
    Generates a beautifully styled Excel (.xlsx) workbook using openpyxl.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]

    # Title Banner
    max_col = max(len(headers), 1)
    end_col_letter = get_column_letter(max_col)
    ws.merge_cells(f'A1:{end_col_letter}1')
    title_cell = ws['A1']
    title_cell.value = f"{Config.COMPANY_NAME} - {title}"
    title_cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 35

    # Headers Row
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append([]) # Empty row 2
    ws.append(headers) # Row 3
    ws.row_dimensions[3].height = 25

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_thin

    # Data Rows
    row_alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    for r_idx, row in enumerate(rows, start=4):
        ws.append(row)
        ws.row_dimensions[r_idx].height = 20
        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=r_idx, column=col_num)
            cell.border = border_thin
            if r_idx % 2 == 0:
                cell.fill = row_alt_fill

    # Auto-fit Column Widths (ignoring merged row 1)
    for col_idx in range(1, len(headers) + 1):
        max_len = len(str(headers[col_idx - 1]))
        for row_idx in range(4, len(rows) + 4):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
